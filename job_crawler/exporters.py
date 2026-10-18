from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from .models import JobItem, clean_text


BASE_HEADERS = [
    "排名", "匹配分", "匹配等级", "推荐动作",
    "平台", "岗位", "公司", "城市", "区域", "薪资", "经验", "学历",
    "技能分", "项目分", "文本相关分", "硬性要求分",
    "命中技能", "缺失技能", "优化建议", "评分说明", "链接",
]


def export_all(
    jobs: list[JobItem],
    output_dir: str | Path,
    *,
    keyword: str,
    city: str,
    meta: dict[str, Any] | None = None,
) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_keyword = _safe_name(keyword)[:30] or "jobs"
    safe_city = _safe_name(city)[:20] or "all"
    base = f"job_matches_{safe_keyword}_{safe_city}_{stamp}"

    files = {
        "json": output / f"{base}.json",
        "csv": output / f"{base}.csv",
        "md": output / f"{base}.md",
    }
    export_json(jobs, files["json"], meta=meta)
    export_csv(jobs, files["csv"])
    export_markdown(jobs, files["md"], keyword=keyword, city=city, meta=meta)

    try:
        xlsx_path = output / f"{base}.xlsx"
        export_xlsx(jobs, xlsx_path)
        files["xlsx"] = xlsx_path
    except Exception as exc:
        # XLSX is optional; CSV/JSON/MD are always available.
        files["xlsx_error"] = Path(str(exc))
    return files


def export_json(jobs: list[JobItem], path: str | Path, meta: dict[str, Any] | None = None) -> None:
    payload = {
        "meta": meta or {},
        "total": len(jobs),
        "jobs": [job.to_dict() for job in jobs],
    }
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def export_csv(jobs: list[JobItem], path: str | Path) -> None:
    with Path(path).open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=BASE_HEADERS)
        writer.writeheader()
        for idx, job in enumerate(jobs, 1):
            writer.writerow(_row(job, idx))


def export_xlsx(jobs: list[JobItem], path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "岗位匹配结果"
    ws.append(BASE_HEADERS)
    for idx, job in enumerate(jobs, 1):
        r = _row(job, idx)
        ws.append([r[h] for h in BASE_HEADERS])

    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 8, "B": 10, "C": 16, "D": 34,
        "E": 12, "F": 28, "G": 24, "H": 10, "I": 14, "J": 14, "K": 12, "L": 10,
        "M": 10, "N": 10, "O": 12, "P": 12,
        "Q": 30, "R": 34, "S": 62, "T": 70, "U": 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def export_markdown(
    jobs: list[JobItem],
    path: str | Path,
    *,
    keyword: str,
    city: str,
    meta: dict[str, Any] | None = None,
) -> None:
    lines: list[str] = []
    lines.append("# 岗位匹配筛选报告")
    lines.append("")
    lines.append(f"- 关键词：{keyword}")
    lines.append(f"- 城市：{city}")
    lines.append(f"- 结果数：{len(jobs)}")
    if meta:
        engine = meta.get("scoring_engine", "")
        if engine:
            lines.append(f"- 评分引擎：{engine}")
        scale = meta.get("score_scale", "")
        if scale:
            lines.append(f"- 评分刻度：{scale}")
        if meta.get("backend_error"):
            lines.append(f"- 评分降级说明：{meta['backend_error']}")
    lines.append("")
    lines.append("> 分数用于岗位优先级排序，不代表录用概率。建议重点看“匹配等级、缺失技能、优化建议”。")
    lines.append("")
    lines.append("## Top 岗位")
    lines.append("")
    lines.append("| 排名 | 分数 | 等级 | 平台 | 岗位 | 公司 | 城市 | 薪资 | 缺失技能 |")
    lines.append("|---:|---:|---|---|---|---|---|---|---|")
    for idx, job in enumerate(jobs[:30], 1):
        lines.append(
            f"| {idx} | {job.total_score:.0f} | {_md(job.match_level)} | {_md(job.platform_name)} | {_md(job.title)} | "
            f"{_md(job.company)} | {_md(job.city or job.district)} | {_md(job.salary)} | "
            f"{_md('、'.join(job.missing_skills[:6]))} |"
        )
    lines.append("")
    lines.append("## 逐岗位建议")
    lines.append("")
    for idx, job in enumerate(jobs, 1):
        lines.append(f"### {idx}. {job.total_score:.0f} 分｜{job.match_level or '-'}｜{job.title} / {job.company} / {job.platform_name}")
        lines.append("")
        if job.url:
            lines.append(f"- 链接：{job.url}")
        lines.append(f"- 推荐动作：{job.recommendation or '-'}")
        lines.append(f"- 地点：{job.location_text() or '-'}")
        lines.append(f"- 薪资：{job.salary or '-'}")
        lines.append(f"- 经验/学历：{' / '.join(x for x in [job.experience, job.degree] if x) or '-'}")
        lines.append(
            f"- 分数拆解：技能 {job.skill_score:.0f} / 项目 {job.project_score:.0f} / "
            f"文本相关 {job.semantic_score:.0f} / 硬性要求 {job.requirement_score:.0f}"
        )
        lines.append(f"- 命中技能：{'、'.join(job.matched_skills) or '-'}")
        lines.append(f"- 缺失技能：{'、'.join(job.missing_skills) or '-'}")
        lines.append(f"- 评分说明：{job.score_reason or '-'}")
        if job.optimization_suggestions:
            lines.append("- 简历优化建议：")
            for sug in job.optimization_suggestions:
                lines.append(f"  - {sug}")
        lines.append("")
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def _row(job: JobItem, idx: int) -> dict[str, Any]:
    return {
        "排名": idx,
        "匹配分": round(float(job.total_score or 0), 1),
        "匹配等级": job.match_level,
        "推荐动作": job.recommendation,
        "平台": job.platform_name,
        "岗位": job.title,
        "公司": job.company,
        "城市": job.city,
        "区域": job.district,
        "薪资": job.salary,
        "经验": job.experience,
        "学历": job.degree,
        "技能分": round(float(job.skill_score or 0), 1),
        "项目分": round(float(job.project_score or 0), 1),
        "文本相关分": round(float(job.semantic_score or 0), 1),
        "硬性要求分": round(float(job.requirement_score or 0), 1),
        "命中技能": "、".join(job.matched_skills),
        "缺失技能": "、".join(job.missing_skills),
        "优化建议": "\n".join(job.optimization_suggestions),
        "评分说明": job.score_reason,
        "链接": job.url,
    }


def _safe_name(text: str) -> str:
    text = clean_text(text)
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in text).strip("_")


def _md(text: Any) -> str:
    return clean_text(text).replace("|", "\\|").replace("\n", " ") or "-"
