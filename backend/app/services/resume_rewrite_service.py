"""Full resume rewrite service.

Uses LLM to rewrite an entire resume tailored to a target JD.
Falls back to rule-based rewriting if LLM is unavailable.
"""

import json
from typing import Optional

from sqlalchemy.orm import Session

from app.models.analysis_task import AnalysisTask, TaskStatus
from app.models.match_report import MatchReport
from app.models.resume import Resume
from app.models.job_description import JobDescription
from app.models.rewritten_resume import RewrittenResume
from app.schemas.rewrite_resume import RewrittenResumeData, ResumeSection, BulletEvidence, EvidenceItem
from app.services.scoring_service import compute_scores
from app.services.diagnosis_service import diagnose
from app.services.skill_normalizer import normalize_skills
from app.core.config import get_settings
from app.core.exceptions import AppException
from app.utils.prompt_render import render_prompt
from app.utils.model_dump import model_to_dict
from app.services.llm_client import is_real_llm_enabled
from app.services.resume_optimization_skill import (
    build_rewrite_guidelines,
    clean_tech_stack,
    contains_inflated_role,
    contains_numeric_claim,
    extract_real_metrics,
    is_resume_style_line,
    parse_json_object_robust,
    project_to_resume_block,
)
from app.utils.export_fonts import safe_export_text, register_pdf_cjk_font, set_docx_default_fonts, set_run_font
from app.services.resume_rewrite_quality import post_process_rewritten_resume


def generate_full_rewrite(task_id: int, db: Session) -> RewrittenResumeData:
    """Generate a full resume rewrite tailored to the target JD."""

    # ── Load task data ───────────────────────────────────
    task = db.get(AnalysisTask, task_id)
    if task is None:
        raise AppException(f"Task {task_id} not found", code="NOT_FOUND", status_code=404)
    if task.status != TaskStatus.completed:
        raise AppException(f"Task {task_id} is not completed", code="NOT_COMPLETED")

    resume = db.get(Resume, task.resume_id)
    jd = db.get(JobDescription, task.jd_id)
    if not resume or not jd:
        raise AppException("Resume or JD not found", code="DATA_MISSING")

    report = db.query(MatchReport).filter(MatchReport.task_id == task_id).first()
    if report is None:
        raise AppException(f"Report for task {task_id} not found", code="NOT_FOUND", status_code=404)

    resume_parsed = json.loads(resume.parsed_json) if resume.parsed_json else {}
    jd_parsed = json.loads(jd.parsed_json) if jd.parsed_json else {}
    matched_skills = json.loads(report.matched_skills_json) if report.matched_skills_json else []
    missing_skills = json.loads(report.missing_skills_json) if report.missing_skills_json else []
    problems_raw = json.loads(report.problems_json) if report.problems_json else []

    # ── Try LLM rewrite ─────────────────────────────────
    rewritten_data = _try_llm_rewrite(
        resume_text=resume.raw_text or "",
        resume_parsed=resume_parsed,
        jd_text=jd.raw_text or "",
        jd_parsed=jd_parsed,
        total_score=report.total_score,
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        problems=problems_raw,
    )

    if rewritten_data is None:
        # Always fall back to rule-based rewrite when LLM is unavailable or fails.
        # LLM_STRICT_MODE only warns the user instead of blocking entirely.
        if is_real_llm_enabled() and get_settings().LLM_STRICT_MODE:
            print("[rewrite] LLM failed but strict mode is on; falling back to rule-based with warning")

        rewritten_data = _rule_based_rewrite(
            resume_parsed=resume_parsed,
            jd_parsed=jd_parsed,
            resume_text=resume.raw_text or "",
            matched_skills=matched_skills,
            missing_skills=missing_skills,
            problems=problems_raw,
        )
        source = "rule"
    else:
        source = "llm"

    # ── Apply universal resume optimization quality gate ──
    # This is where the general resume-polishing process is embedded into
    # the product: clean tech-stack lines, enforce intern-appropriate tone,
    # flag unverified metrics/overclaims, and expose rewrite_strategy /
    # quality_gate to the frontend/API response.
    rewritten_data = post_process_rewritten_resume(
        rewritten_data,
        resume_text=resume.raw_text or "",
        resume_parsed=resume_parsed,
        jd_parsed=jd_parsed,
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        problems=problems_raw,
    )

    # ── Persist ──────────────────────────────────────────
    existing = db.query(RewrittenResume).filter(RewrittenResume.task_id == task_id).first()
    if existing:
        existing.rewritten_text = rewritten_data.full_text
        existing.rewritten_json = json.dumps(model_to_dict(rewritten_data), ensure_ascii=False)
        existing.source = source
        db.commit()
        db.refresh(existing)
    else:
        record = RewrittenResume(
            task_id=task_id,
            resume_id=task.resume_id,
            jd_id=task.jd_id,
            rewritten_text=rewritten_data.full_text,
            rewritten_json=json.dumps(model_to_dict(rewritten_data), ensure_ascii=False),
            source=source,
        )
        db.add(record)
        db.commit()

    return rewritten_data


def get_rewrite(task_id: int, db: Session) -> Optional[RewrittenResumeData]:
    """Get existing rewrite for a task."""
    record = db.query(RewrittenResume).filter(RewrittenResume.task_id == task_id).first()
    if record is None:
        return None

    if record.rewritten_json:
        data = json.loads(record.rewritten_json)
        if record.edited_text:
            data["full_text"] = record.edited_text
        return RewrittenResumeData(**data)
    if record.edited_text or record.rewritten_text:
        text = record.edited_text or record.rewritten_text
        return RewrittenResumeData(
            job_title="",
            target_company="",
            summary="",
            sections=[],
            highlights=[],
            full_text=text,
            changes_summary="已加载用户保存的改写版本。",
        )
    return None


def save_edited_rewrite(task_id: int, edited_text: str, db: Session) -> None:
    """Save user-edited version of the rewritten resume."""
    record = db.query(RewrittenResume).filter(RewrittenResume.task_id == task_id).first()
    if record is None:
        raise AppException(f"Rewrite for task {task_id} not found", code="NOT_FOUND", status_code=404)
    record.edited_text = edited_text
    db.commit()


# ── LLM rewrite ─────────────────────────────────────────

def _try_llm_rewrite(
    resume_text: str,
    resume_parsed: dict,
    jd_text: str,
    jd_parsed: dict,
    total_score: float,
    matched_skills: list[str],
    missing_skills: list[str],
    problems: list[dict],
) -> Optional[RewrittenResumeData]:
    """Try to generate rewrite via LLM. Returns None if LLM unavailable."""
    try:
        from app.services.llm_client import get_llm_client
        from app.prompts import FULL_REWRITE_SYSTEM_PROMPT, FULL_REWRITE_USER_TEMPLATE

        client = get_llm_client()
        if type(client).__name__ == 'MockLLMClient':
            return None

        # Build match level
        if total_score >= 85:
            match_level = "非常匹配"
        elif total_score >= 70:
            match_level = "较匹配"
        elif total_score >= 50:
            match_level = "一般"
        else:
            match_level = "较不匹配"

        # Build problems summary
        problem_texts = []
        for p in problems:
            if isinstance(p, dict):
                problem_texts.append(f"[{p.get('severity', '')}] {p.get('title', '')}: {p.get('suggestion', '')}")
            else:
                problem_texts.append(str(p))

        user_prompt = render_prompt(
            FULL_REWRITE_USER_TEMPLATE,
            resume_text=resume_text[:3000],
            resume_parsed=json.dumps(resume_parsed, ensure_ascii=False, indent=2)[:2000],
            jd_text=jd_text[:2000],
            jd_parsed=json.dumps(jd_parsed, ensure_ascii=False, indent=2)[:2000],
            total_score=total_score,
            match_level=match_level,
            matched_skills=", ".join(matched_skills) or "无",
            missing_skills=", ".join(missing_skills) or "无",
            problems="\n".join(problem_texts) or "无",
            rewrite_guidelines=build_rewrite_guidelines(),
        )

        raw = client.chat(FULL_REWRITE_SYSTEM_PROMPT, user_prompt)

        # Parse JSON response with repairs for common LLM issues:
        # fenced code blocks, surrounding prose and trailing commas.
        data = parse_json_object_robust(raw)
        if data is None:
            return None

        # Validate and convert
        sections = []
        for s in data.get("sections", []):
            if isinstance(s, dict):
                sections.append(ResumeSection(
                    title=s.get("title", ""),
                    content=s.get("content", ""),
                ))

        llm_summary = data.get("summary", "")

        # Parse evidence chain from LLM response (v2: split resume/jd evidence)
        evidence_chain: list[BulletEvidence] = []
        for item in data.get("evidence_chain", []):
            if not isinstance(item, dict):
                continue

            resume_ev = _parse_evidence_list(item.get("resume_evidence", []))
            jd_ev = _parse_evidence_list(item.get("jd_evidence", []))

            # If LLM returned old format (flat "evidence" list), split by source_type
            if not resume_ev and not jd_ev:
                flat = _parse_evidence_list(item.get("evidence", []))
                for e in flat:
                    if e.source_type.startswith("resume_"):
                        resume_ev.append(e)
                    else:
                        jd_ev.append(e)

            # Do not trust LLM-provided risk blindly. Recompute risk from
            # split evidence so JD evidence alone never lowers fabrication risk.
            risk, fab, need_confirm = _assess_split_evidence(
                resume_ev,
                jd_ev,
                str(item.get("rewritten_bullet", "")),
            )
            # Preserve an explicit human-review flag from LLM when it is stricter.
            need_confirm = need_confirm or bool(item.get("need_manual_confirm", False))

            evidence_chain.append(BulletEvidence(
                rewritten_bullet=str(item.get("rewritten_bullet", "")),
                resume_evidence=resume_ev,
                jd_evidence=jd_ev,
                risk=risk,
                fabrication_check=fab,
                need_manual_confirm=need_confirm,
            ))

        # If LLM didn't provide evidence chain, generate it from rule-based analysis
        if not evidence_chain:
            evidence_chain = _generate_evidence_chain(
                resume_parsed=resume_parsed,
                jd_parsed=jd_parsed,
                matched_skills=matched_skills,
                missing_skills=missing_skills,
                sections=sections,
                summary=llm_summary,
            )

        result = RewrittenResumeData(
            job_title=data.get("job_title", ""),
            target_company=data.get("target_company", ""),
            summary=llm_summary,
            sections=sections,
            highlights=data.get("highlights", []),
            full_text=data.get("full_text", ""),
            changes_summary=data.get("changes_summary", ""),
            evidence_chain=evidence_chain,
        )

        # Validate that the LLM actually produced content, not just empty placeholders
        if not result.full_text.strip() and not result.summary.strip() and len(sections) == 0:
            print("[rewrite] LLM returned empty response, will fall back to rule-based")
            return None

        return result

    except Exception as e:
        print(f"[rewrite] LLM rewrite failed, falling back to rule-based: {e}")
        return None


# ── Rule-based rewrite ───────────────────────────────────

def _rule_based_rewrite(
    resume_parsed: dict,
    jd_parsed: dict,
    resume_text: str,
    matched_skills: list[str],
    missing_skills: list[str],
    problems: list[dict],
) -> RewrittenResumeData:
    """Generate a rewritten resume using rule-based approach.

    Uses resume_parsed primarily, but falls back to resume_text for sections
    that are missing from the structured parse so the output is never empty.
    """

    # Enrich resume_parsed from resume_text for missing sections
    enriched = dict(resume_parsed)
    _enrich_from_text(enriched, resume_text)

    jd_title = jd_parsed.get("title", "目标岗位")
    jd_company = jd_parsed.get("company", "")

    # Build professional summary — always populated from enriched data
    summary = _build_summary(enriched, jd_parsed, matched_skills)

    sections: list[ResumeSection] = []

    # Skills section — always present
    skills_content = _rewrite_skills_section(enriched, jd_parsed, matched_skills, missing_skills)
    sections.append(ResumeSection(title="专业技能", content=skills_content))

    # Work experience section
    work_content = _rewrite_work_section(enriched, jd_parsed)
    if work_content:
        sections.append(ResumeSection(title="工作经历", content=work_content))

    # Projects section
    project_content = _rewrite_projects_section(enriched, jd_parsed, matched_skills)
    if project_content:
        sections.append(ResumeSection(title="项目经历", content=project_content))

    # Education section
    edu_content = _rewrite_education_section(enriched)
    if edu_content:
        sections.append(ResumeSection(title="教育背景", content=edu_content))

    # Other section
    other_content = _rewrite_other_section(enriched)
    if other_content:
        sections.append(ResumeSection(title="其他", content=other_content))

    # Ensure we always have at least a minimal structure
    if len(sections) <= 1:
        sections.append(_build_generic_section(enriched, resume_text, jd_parsed))

    # Build highlights
    highlights: list[str] = []
    if matched_skills:
        highlights.append(f"突出了与岗位匹配的 {len(matched_skills)} 项核心技能")
    if missing_skills:
        highlights.append(f"标注了 {len(missing_skills)} 项待补充技能")
    if resume_text.strip():
        highlights.append("基于原始简历文本进行了重组优化")
    else:
        highlights.append("请上传简历以获取个性化优化")

    # Build full text
    full_text = _assemble_full_text(jd_title, jd_company, summary, sections, missing_skills)

    # Generate evidence chain from rule-based analysis
    evidence_chain = _generate_evidence_chain(
        resume_parsed=enriched,
        jd_parsed=jd_parsed,
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        sections=sections,
        summary=summary,
    )

    result = RewrittenResumeData(
        job_title=jd_title,
        target_company=jd_company,
        summary=summary,
        sections=sections,
        highlights=highlights,
        full_text=full_text,
        changes_summary="基于规则的简历优化：重组技能展示顺序、优化项目描述、补充缺失技能标注。",
        evidence_chain=evidence_chain,
    )
    return post_process_rewritten_resume(
        result,
        resume_text=resume_text,
        resume_parsed=enriched,
        jd_parsed=jd_parsed,
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        problems=problems,
    )


def _enrich_from_text(parsed: dict, resume_text: str) -> None:
    """Fill missing parsed fields from resume_text when the structured parse is sparse."""
    text = (resume_text or "").strip()
    if not text:
        return

    # Extract name from first line if missing
    if not parsed.get("name"):
        lines = text.splitlines()
        if lines:
            first = lines[0].strip()
            if len(first) <= 20 and not any(k in first for k in ["简历", "个人", "求职", "联系"]):
                parsed["name"] = first

    # Extract skills from text if missing
    if not parsed.get("skills"):
        from app.services.skill_normalizer import normalize_skills
        from app.services.platform_crawler_service import _extract_skills_from_text
        skills = _extract_skills_from_text(text)
        if skills:
            parsed["skills"] = skills

    # Extract education keywords
    if not parsed.get("education"):
        import re
        edu_lines: list[str] = []
        for m in re.finditer(r"(?:学校|学院|大学|教育|学历|本科|硕士|博士|大专|毕业|专业)[\s：:]*([^\n]{2,60})", text):
            edu_lines.append(m.group(0).strip())
        if edu_lines:
            parsed["education"] = [{"raw": line} for line in edu_lines[:3]]

    # Extract experience hints
    if not parsed.get("work_experience"):
        import re
        exp_hints: list[dict] = []
        for m in re.finditer(r"(?:工作经历|实习经历|工作经验|在职|任职)[\s：:]*([^\n]{2,200})", text):
            exp_hints.append({"description": m.group(0).strip()})
        if exp_hints:
            parsed["work_experience"] = exp_hints


def _build_generic_section(parsed: dict, resume_text: str, jd_parsed: dict) -> ResumeSection:
    """Build a generic content section from raw text when structured data is sparse."""
    lines = (resume_text or "").strip().splitlines()
    content_lines = [line.strip() for line in lines if line.strip()]
    # Skip obvious header lines
    skip_keywords = ["简历", "个人简历", "联系方式", "求职意向"]
    filtered = [line for line in content_lines if not any(k in line for k in skip_keywords)]
    if not filtered:
        jd_title = jd_parsed.get("title", "目标岗位")
        return ResumeSection(
            title="简历正文",
            content=f"[请补充：个人经历、项目描述、技能等内容，以匹配 {jd_title} 岗位要求]",
        )
    return ResumeSection(title="简历正文", content="\n".join(filtered[:30]))


def _build_summary(resume_parsed: dict, jd_parsed: dict, matched_skills: list[str]) -> str:
    """Build a concise intern-friendly summary without overclaiming."""
    basic = resume_parsed.get("basic_info", {}) if isinstance(resume_parsed.get("basic_info"), dict) else {}
    name = resume_parsed.get("name") or basic.get("name", "")
    jd_title = jd_parsed.get("job_title") or jd_parsed.get("title", "目标岗位")
    skills = matched_skills[:5]

    subject = name or "候选人"
    if skills:
        skill_text = "、".join(skills)
        return (
            f"{subject}具备与{jd_title}相关的项目实践，已在简历中体现{skill_text}等能力；"
            "关注可验证的项目证据、后端接口实现与大模型应用落地，适合以实习生身份参与需求拆解、功能开发和迭代优化。"
        )
    return (
        f"{subject}具备计算机与人工智能相关学习背景，正在补强与{jd_title}相关的项目实践；"
        "简历改写将优先保留真实经历，并对缺失能力使用待补充提示。"
    )


def _rewrite_skills_section(
    resume_parsed: dict,
    jd_parsed: dict,
    matched: list[str],
    missing: list[str],
) -> str:
    """Rewrite skills section with evidence-safe grouping.

    JD-only skills are never written as the candidate's existing skills.
    They are shown as "待补充" instead.
    """
    lines: list[str] = []

    existing_skills = resume_parsed.get("skills", [])
    if isinstance(existing_skills, str):
        existing_skills = [s.strip() for s in existing_skills.split(",") if s.strip()]
    existing_skills = clean_tech_stack(existing_skills, max_items=18)

    matched_set = {s.lower() for s in matched}
    matched_items = [s for s in existing_skills if s.lower() in matched_set]
    other_items = [s for s in existing_skills if s.lower() not in matched_set]

    if matched_items:
        lines.append(f"岗位相关技能：{'、'.join(matched_items[:10])}")
    if other_items:
        lines.append(f"其他已掌握技能：{'、'.join(other_items[:10])}")
    if matched and not matched_items:
        # Matched skills may come from projects rather than the explicit skill section.
        lines.append(f"项目中已体现能力：{'、'.join(matched[:10])}")
    if missing:
        lines.append(f"待补充技能：{'、'.join(missing[:6])}（仅在学习并实践后写入正式简历）")

    if not lines:
        jd_title = jd_parsed.get("job_title") or jd_parsed.get("title", "目标岗位")
        lines.append(f"[请补充：与 {jd_title} 相关且已有项目证据支撑的技能]")

    return "\n".join(lines)


def _rewrite_work_section(resume_parsed: dict, jd_parsed: dict) -> str:
    """Rewrite work experience section."""
    experiences = resume_parsed.get("work_experience", [])
    if not experiences:
        return ""

    lines = []
    for exp in experiences:
        if isinstance(exp, dict):
            company = exp.get("company", "")
            title = exp.get("title", exp.get("position", ""))
            period = exp.get("period", exp.get("duration", ""))
            desc = exp.get("description", exp.get("responsibilities", ""))

            header = f"{title}"
            if company:
                header += f" | {company}"
            if period:
                header += f" ({period})"
            lines.append(header)

            if desc:
                if isinstance(desc, list):
                    for d in desc:
                        lines.append(f"  - {d}")
                else:
                    lines.append(f"  - {desc}")
        elif isinstance(exp, str):
            lines.append(exp)

    return "\n".join(lines)


def _rewrite_projects_section(
    resume_parsed: dict,
    jd_parsed: dict,
    matched: list[str],
) -> str:
    """Rewrite projects in resume-friendly structure.

    Format follows the reusable resume-polishing workflow:
    项目简介 → 技术栈 → 技术亮点. It avoids turning the project into a long
    narrative and keeps claims evidence-based.
    """
    projects = resume_parsed.get("projects", [])
    if not projects:
        return ""

    blocks: list[str] = []
    for proj in projects:
        if isinstance(proj, dict):
            blocks.append(project_to_resume_block(proj, matched_skills=matched))
        elif isinstance(proj, str) and proj.strip():
            blocks.append(proj.strip())

    return "\n\n".join(blocks).strip()


def _rewrite_education_section(resume_parsed: dict) -> str:
    """Rewrite education section."""
    edu = resume_parsed.get("education", [])
    if not edu:
        return ""

    lines = []
    for e in edu:
        if isinstance(e, dict):
            school = e.get("school", e.get("institution", ""))
            degree = e.get("degree", "")
            major = e.get("major", e.get("field", ""))
            period = e.get("period", e.get("duration", ""))

            parts = [p for p in [school, degree, major] if p]
            line = " | ".join(parts)
            if period:
                line += f" ({period})"
            lines.append(line)
        elif isinstance(e, str):
            lines.append(e)

    return "\n".join(lines)


def _rewrite_other_section(resume_parsed: dict) -> str:
    """Rewrite other/misc section."""
    parts = []

    certs = resume_parsed.get("certifications", [])
    if certs:
        if isinstance(certs, list):
            parts.append(f"证书: {'、'.join(str(c) for c in certs)}")
        else:
            parts.append(f"证书: {certs}")

    awards = resume_parsed.get("awards", [])
    if awards:
        if isinstance(awards, list):
            parts.append(f"荣誉: {'、'.join(str(a) for a in awards)}")
        else:
            parts.append(f"荣誉: {awards}")

    return "\n".join(parts) if parts else ""


def _assemble_full_text(
    jd_title: str,
    jd_company: str,
    summary: str,
    sections: list[ResumeSection],
    missing: list[str],
) -> str:
    """Assemble full resume text."""
    lines = []
    lines.append(f"目标岗位: {jd_title}")
    if jd_company:
        lines.append(f"目标公司: {jd_company}")
    lines.append("")
    lines.append("个人简介")
    lines.append(summary)
    lines.append("")

    for sec in sections:
        lines.append(sec.title)
        lines.append(sec.content)
        lines.append("")

    if missing:
        lines.append("待补充技能")
        lines.append(f"以下技能是目标岗位要求但简历中尚未体现的，建议学习后补充: {', '.join(missing)}")

    return "\n".join(lines)


# ── Evidence chain generation ────────────────────────────

def _parse_evidence_list(raw: list) -> list["EvidenceItem"]:
    """Parse a list of evidence dicts from LLM or rule-based sources."""
    items: list[EvidenceItem] = []
    valid_types = {
        "resume_project", "resume_skill", "resume_work_experience",
        "resume_education", "jd_requirement", "jd_preferred", "system_rule",
    }
    for ev in (raw or []):
        if not isinstance(ev, dict):
            continue
        st = str(ev.get("source_type", "system_rule"))
        if st not in valid_types:
            src = str(ev.get("source", ""))
            if src.startswith("project."):
                st = "resume_project"
            elif src.startswith("resume.skills"):
                st = "resume_skill"
            elif src.startswith("resume.work"):
                st = "resume_work_experience"
            elif src.startswith("resume.education"):
                st = "resume_education"
            elif src.startswith("jd.preferred"):
                st = "jd_preferred"
            elif src.startswith("jd."):
                st = "jd_requirement"
            else:
                st = "system_rule"
        items.append(EvidenceItem(
            source=str(ev.get("source", "")),
            source_type=st,
            field=str(ev.get("field", "")),
            confidence=float(ev.get("confidence", 0.0)),
        ))
    return items


def _generate_evidence_chain(
    resume_parsed: dict,
    jd_parsed: dict,
    matched_skills: list[str],
    missing_skills: list[str],
    sections: list[ResumeSection],
    summary: str,
) -> list[BulletEvidence]:
    """Generate evidence chain with resume/jd evidence split.

    Resume evidence = proof the candidate did something.
    JD evidence = what the job requires (direction, not proof).
    Risk is driven primarily by resume evidence quality.
    """
    evidence_chain: list[BulletEvidence] = []

    resume_projects = resume_parsed.get("projects", []) or []
    resume_skills = resume_parsed.get("skills", []) or []
    resume_work = resume_parsed.get("work_experience", []) or []
    resume_edu = resume_parsed.get("education", []) or []
    jd_required = jd_parsed.get("required_skills", []) or []
    jd_preferred = jd_parsed.get("preferred_skills", []) or []
    jd_title = jd_parsed.get("title", "")

    seen_bullets: set[str] = set()
    all_lines: list[str] = []

    for section in sections:
        content = section.content
        if not content.strip():
            continue
        for line in content.split("\n"):
            line = line.strip()
            if line and len(line) >= 10:
                all_lines.append(line)

    if summary.strip():
        for line in summary.split("。"):
            line = line.strip()
            if line and len(line) >= 10:
                all_lines.append(line)

    for line in all_lines:
        if line in seen_bullets:
            continue
        seen_bullets.add(line)

        resume_ev, jd_ev = _find_split_evidence(
            line=line,
            resume_projects=resume_projects,
            resume_skills=resume_skills,
            resume_work=resume_work,
            resume_edu=resume_edu,
            jd_required=jd_required,
            jd_preferred=jd_preferred,
            jd_title=jd_title,
        )

        risk, fab, need_confirm = _assess_split_evidence(resume_ev, jd_ev, line)

        evidence_chain.append(BulletEvidence(
            rewritten_bullet=line,
            resume_evidence=resume_ev,
            jd_evidence=jd_ev,
            risk=risk,
            fabrication_check=fab,
            need_manual_confirm=need_confirm,
        ))

    return evidence_chain


def _find_split_evidence(
    line: str,
    resume_projects: list[dict],
    resume_skills: list,
    resume_work: list[dict],
    resume_edu: list[dict],
    jd_required: list,
    jd_preferred: list,
    jd_title: str,
) -> tuple[list[EvidenceItem], list[EvidenceItem]]:
    """Find resume evidence and JD evidence separately for a bullet."""
    resume_ev: list[EvidenceItem] = []
    jd_ev: list[EvidenceItem] = []
    line_lower = line.lower()

    # ── Resume evidence (proof of experience) ──

    for proj in resume_projects:
        if not isinstance(proj, dict):
            continue
        proj_name = str(proj.get("name", ""))
        tech_stack = proj.get("tech_stack", [])
        if isinstance(tech_stack, str):
            tech_stack = [tech_stack]
        desc = str(proj.get("description", ""))

        matched_tech = [str(t) for t in tech_stack if str(t).lower() in line_lower]
        if matched_tech:
            resume_ev.append(EvidenceItem(
                source=f"project.{proj_name}.tech_stack",
                source_type="resume_project",
                field=", ".join(matched_tech[:5]),
                confidence=0.90,
            ))

        if proj_name and proj_name.lower() in line_lower:
            resume_ev.append(EvidenceItem(
                source=f"project.{proj_name}",
                source_type="resume_project",
                field=proj_name,
                confidence=0.85,
            ))

        if desc:
            desc_words = set(desc.lower().split())
            line_words = set(line_lower.split())
            overlap = desc_words & line_words
            if len(overlap) >= 4:
                resume_ev.append(EvidenceItem(
                    source=f"project.{proj_name}.description",
                    source_type="resume_project",
                    field="关键词匹配: " + ", ".join(sorted(overlap)[:5]),
                    confidence=0.75,
                ))

    for skill in resume_skills:
        skill_str = str(skill).lower() if isinstance(skill, str) else str(
            skill.get("name", skill.get("raw", ""))).lower()
        if skill_str and len(skill_str) >= 2 and skill_str in line_lower:
            label = str(skill) if isinstance(skill, str) else skill.get("name", skill.get("raw", ""))
            resume_ev.append(EvidenceItem(
                source="resume.skills",
                source_type="resume_skill",
                field=label,
                confidence=0.80,
            ))

    for exp in resume_work:
        if not isinstance(exp, dict):
            continue
        company = str(exp.get("company", ""))
        title = str(exp.get("title", exp.get("position", "")))
        if company and company.lower() in line_lower:
            resume_ev.append(EvidenceItem(
                source="resume.work_experience",
                source_type="resume_work_experience",
                field=company,
                confidence=0.85,
            ))
        if title and title.lower() in line_lower:
            resume_ev.append(EvidenceItem(
                source="resume.work_experience",
                source_type="resume_work_experience",
                field=title,
                confidence=0.85,
            ))

    for edu in resume_edu:
        if not isinstance(edu, dict):
            continue
        school = str(edu.get("school", edu.get("institution", "")))
        degree = str(edu.get("degree", ""))
        if school and school.lower() in line_lower:
            resume_ev.append(EvidenceItem(
                source="resume.education",
                source_type="resume_education",
                field=school,
                confidence=0.85,
            ))
        if degree and degree.lower() in line_lower:
            resume_ev.append(EvidenceItem(
                source="resume.education",
                source_type="resume_education",
                field=degree,
                confidence=0.85,
            ))

    # ── JD evidence (job requirements — direction only) ──

    for skill in jd_required:
        skill_str = str(skill).lower() if isinstance(skill, str) else str(skill.get("name", "")).lower()
        if skill_str and len(skill_str) >= 2 and skill_str in line_lower:
            jd_ev.append(EvidenceItem(
                source="jd.requirement",
                source_type="jd_requirement",
                field=str(skill) if isinstance(skill, str) else skill.get("name", ""),
                confidence=0.70,
            ))

    for skill in jd_preferred:
        skill_str = str(skill).lower() if isinstance(skill, str) else str(skill.get("name", "")).lower()
        if skill_str and len(skill_str) >= 2 and skill_str in line_lower:
            jd_ev.append(EvidenceItem(
                source="jd.preferred",
                source_type="jd_preferred",
                field=str(skill) if isinstance(skill, str) else skill.get("name", ""),
                confidence=0.60,
            ))

    if jd_title and jd_title.lower() in line_lower:
        jd_ev.append(EvidenceItem(
            source="jd.title",
            source_type="jd_requirement",
            field=jd_title,
            confidence=0.65,
        ))

    return resume_ev, jd_ev


def _assess_split_evidence(
    resume_evidence: list[EvidenceItem],
    jd_evidence: list[EvidenceItem],
    bullet: str,
) -> tuple[str, str, bool]:
    """Assess rewrite risk using the resume-optimization quality gate.

    Core rule: JD evidence only tells us what the job wants; it never proves the
    candidate has done it. Resume evidence is the only proof source.
    """
    has_resume = len(resume_evidence) > 0

    if not has_resume:
        return ("high", "warning", True)

    # Unsupported numeric claims are high-risk unless the same number appears in
    # resume-side evidence fields. This prevents invented "提升30%" style claims.
    if contains_numeric_claim(bullet):
        bullet_metrics = set(extract_real_metrics(bullet))
        evidence_text = " ".join(e.field for e in resume_evidence)
        evidence_metrics = set(extract_real_metrics(evidence_text))
        if bullet_metrics and not (bullet_metrics & evidence_metrics):
            return ("high", "warning", True)

    # Role inflation should be manually checked unless backed by project evidence.
    if contains_inflated_role(bullet):
        has_project_strong = any(e.source_type == "resume_project" and e.confidence >= 0.85 for e in resume_evidence)
        if not has_project_strong:
            return ("high", "warning", True)

    if not is_resume_style_line(bullet):
        return ("medium", "warning", True)

    has_project_evidence = any(e.source_type == "resume_project" for e in resume_evidence)
    has_work_evidence = any(e.source_type == "resume_work_experience" for e in resume_evidence)
    resume_confs = [e.confidence for e in resume_evidence]
    avg_resume = sum(resume_confs) / len(resume_confs)
    has_strong = any(e.confidence >= 0.80 for e in resume_evidence)

    if has_project_evidence and has_strong:
        return ("low", "passed", False)
    if has_work_evidence and has_strong:
        return ("medium", "passed", False)
    if has_strong or avg_resume >= 0.75:
        # Skill-only evidence is real but weaker than project evidence, so require review.
        return ("medium", "passed", True)
    return ("medium", "passed", True)

# ── Export rewritten resume ─────────────────────────────

def get_effective_rewrite_text(task_id: int, db: Session) -> str:
    record = db.query(RewrittenResume).filter(RewrittenResume.task_id == task_id).first()
    if record is None:
        raise AppException(f"Rewrite for task {task_id} not found", code="NOT_FOUND", status_code=404)
    text = record.edited_text or record.rewritten_text or ""
    if not text.strip():
        raise AppException("改写简历内容为空，无法导出", code="EMPTY_REWRITE")
    return text


def _split_resume_sections(text: str) -> list[tuple[str, list[str]]]:
    section_names = [
        "个人优势", "个人简介", "求职意向", "专业技能", "技能", "项目经历", "工作经历",
        "实习经历", "教育背景", "教育经历", "荣誉奖项", "职业证书", "证书", "自我评价", "其他"
    ]
    sections: list[tuple[str, list[str]]] = []
    current_title = "个人信息"
    current_lines: list[str] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        normalized = line.rstrip("：:")
        is_section = normalized in section_names or (len(normalized) <= 8 and normalized.endswith(("经历", "背景", "技能", "证书", "奖项", "优势")))
        if is_section:
            if current_lines:
                sections.append((current_title, current_lines))
            current_title = normalized
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines:
        sections.append((current_title, current_lines))
    return sections or [("简历", [text])]


def _guess_resume_name_and_contact(text: str) -> tuple[str, str]:
    lines = [x.strip() for x in (text or "").splitlines() if x.strip()]
    name = "改写简历"
    contact_parts = []
    if lines:
        first = lines[0]
        if len(first) <= 40 and not any(k in first for k in ["目标岗位", "求职意向", "个人", "专业技能"]):
            name = first.split("|")[0].strip() or name
    import re
    phone = re.search(r"1[3-9]\d{9}", text or "")
    email = re.search(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", text or "")
    if phone:
        contact_parts.append(phone.group(0))
    if email:
        contact_parts.append(email.group(0))
    return name, " | ".join(contact_parts)



def export_rewritten_resume_docx(task_id: int, db: Session) -> bytes:
    """Export rewritten resume as a styled Word document, not plain text."""
    text = safe_export_text(get_effective_rewrite_text(task_id, db))
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise AppException("缺少 python-docx，无法导出 Word。请安装 requirements.txt。", code="DOCX_EXPORT_FAILED") from exc

    from io import BytesIO

    name, contact = _guess_resume_name_and_contact(text)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    set_docx_default_fonts(doc, east_asia="Microsoft YaHei", latin="Arial")

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    run = title_p.add_run(safe_export_text(name or "候选人"))
    set_run_font(run, size=18, bold=True)
    run.font.color.rgb = RGBColor(31, 78, 121)

    if contact:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(safe_export_text(contact))
        set_run_font(r, size=9)
        r.font.color.rgb = RGBColor(90, 90, 90)

    def add_section_title(title: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(safe_export_text(title))
        set_run_font(r, size=11.2, bold=True)
        r.font.color.rgb = RGBColor(31, 78, 121)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '4F81BD')
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    def add_line(line: str) -> None:
        line = safe_export_text(line).strip()
        if not line:
            return
        bullet = line.startswith(('-', '•', '·', '*'))
        p = doc.add_paragraph(style=None)
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.line_spacing = 1.05
        if bullet:
            p.paragraph_format.left_indent = Inches(0.18)
            text_line = line.lstrip('-•·* ').strip()
            r0 = p.add_run('• ')
            set_run_font(r0, size=9.5)
            r0.font.color.rgb = RGBColor(31, 78, 121)
        else:
            text_line = line
        # Bold short labels before Chinese colon.
        if '：' in text_line and len(text_line.split('：', 1)[0]) <= 12:
            left, right = text_line.split('：', 1)
            r1 = p.add_run(left + '：')
            set_run_font(r1, size=9.8, bold=True)
            r2 = p.add_run(right)
            set_run_font(r2, size=9.8)
        else:
            r = p.add_run(text_line)
            set_run_font(r, size=9.8)

    for title, lines in _split_resume_sections(text):
        if title == "个人信息":
            filtered = [x for x in lines if (name and name not in x) and (not contact or x not in contact)]
            if filtered:
                add_section_title("个人信息")
                for line in filtered[:3]:
                    add_line(line)
            continue
        add_section_title(title)
        for line in lines:
            add_line(line)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def export_rewritten_resume_pdf(task_id: int, db: Session) -> bytes:
    """Export the edited/rewritten resume as a styled PDF using ReportLab CJK font."""
    text = safe_export_text(get_effective_rewrite_text(task_id, db))
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    except ImportError as exc:
        raise AppException("缺少 reportlab，无法导出 PDF。请安装 requirements.txt。", code="PDF_EXPORT_FAILED") from exc

    from io import BytesIO
    import html

    font = register_pdf_cjk_font()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CNTitle", fontName=font, fontSize=18, leading=22, alignment=1, textColor=colors.HexColor("#1f4e79"), spaceAfter=4))
    styles.add(ParagraphStyle(name="CNContact", fontName=font, fontSize=8.8, leading=12, alignment=1, textColor=colors.HexColor("#555555"), spaceAfter=6))
    styles.add(ParagraphStyle(name="CNSection", fontName=font, fontSize=11.2, leading=14, textColor=colors.HexColor("#1f4e79"), spaceBefore=5, spaceAfter=2))
    styles.add(ParagraphStyle(name="CNBody", fontName=font, fontSize=9.2, leading=12.5, spaceAfter=1.6))
    styles.add(ParagraphStyle(name="CNBullet", fontName=font, fontSize=9.0, leading=12.2, leftIndent=10, firstLineIndent=-8, spaceAfter=1.2))

    name, contact = _guess_resume_name_and_contact(text)
    esc = lambda s: html.escape(safe_export_text(s or ""))
    story = [Paragraph(esc(name or "候选人"), styles["CNTitle"])]
    if contact:
        story.append(Paragraph(esc(contact), styles["CNContact"]))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#4F81BD")))
    story.append(Spacer(1, 3))

    for title, lines in _split_resume_sections(text):
        filtered = lines if title != "个人信息" else [x for x in lines if name not in x]
        if not filtered:
            continue
        story.append(Paragraph(esc(title), styles["CNSection"]))
        story.append(HRFlowable(width="100%", thickness=0.45, color=colors.HexColor("#D9E2F3")))
        for line in filtered:
            line = safe_export_text(line).strip()
            if not line:
                continue
            stripped = line.lstrip('-•·* ').strip()
            if line.startswith(('-', '•', '·', '*')):
                story.append(Paragraph("- " + esc(stripped), styles["CNBullet"]))
            else:
                story.append(Paragraph(esc(stripped), styles["CNBody"]))
        story.append(Spacer(1, 1.5))

    doc.build(story)
    return buf.getvalue()
