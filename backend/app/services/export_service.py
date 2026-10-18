"""Export analysis report as Word/PDF with Chinese-safe fonts."""

from __future__ import annotations

import io
import json
import html
from typing import Any, Tuple

from sqlalchemy.orm import Session

from app.models.analysis_task import AnalysisTask
from app.models.match_report import MatchReport
from app.models.resume import Resume
from app.models.job_description import JobDescription
from app.schemas.analysis import ScoreDetail
from app.core.exceptions import AppException
from app.utils.export_fonts import safe_export_text, register_pdf_cjk_font, set_docx_default_fonts, set_run_font


def export_report_docx(task_id: int, db: Session) -> bytes:
    """Generate a readable Word document for the analysis report."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except Exception as exc:
        raise AppException("缺少 python-docx，无法导出 Word。请安装 requirements.txt。", code="DOCX_EXPORT_FAILED") from exc

    task, report, resume, jd = _load_data(task_id, db)
    score_detail, ai = _load_score_detail(report)
    problems = _load_list(report.problems_json)
    matched = _load_list(report.matched_skills_json)
    missing = _load_list(report.missing_skills_json)
    suggestions = _load_list(report.suggestions_json)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    set_docx_default_fonts(doc)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI 简历 JD 匹配分析报告")
    set_run_font(r, size=18, bold=True)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"简历：{resume.filename or ''}    岗位：{jd.title or ''}    公司：{jd.company or ''}")
    set_run_font(r, size=9)
    r.font.color.rgb = RGBColor(90, 90, 90)

    _add_score_table(doc, report.total_score or 0, score_detail)

    _add_heading(doc, "综合结论")
    conclusion = (ai.get("overall_conclusion") or _build_basic_conclusion(report.total_score or 0, matched, missing))
    _add_body(doc, conclusion)

    _add_heading(doc, "AI 深度分析")
    if ai:
        for label, key in [
            ("匹配定位", "match_positioning"),
            ("简历改写重点", "resume_rewrite_strategy"),
            ("面试风险", "interview_risks"),
            ("建议行动", "action_plan"),
        ]:
            value = ai.get(key)
            if value:
                _add_subheading(doc, label)
                if isinstance(value, list):
                    for item in value[:8]:
                        text = item.get("action") or item.get("reason") if isinstance(item, dict) else str(item)
                        _add_bullet(doc, text)
                else:
                    _add_body(doc, str(value))
    else:
        _add_body(doc, "本报告基于已保存的匹配分数、技能命中和简历问题生成。")

    _add_heading(doc, "技能匹配")
    _add_body(doc, "命中技能：" + ("、".join(map(str, matched)) if matched else "暂无"))
    _add_body(doc, "缺失/风险技能：" + ("、".join(map(str, missing)) if missing else "暂无"))

    _add_heading(doc, "简历问题与优化建议")
    if problems:
        for item in problems:
            if not isinstance(item, dict):
                continue
            _add_subheading(doc, item.get("title") or "简历问题")
            _add_body(doc, "证据：" + str(item.get("evidence") or ""))
            _add_body(doc, "建议：" + str(item.get("suggestion") or ""))
    elif suggestions:
        for item in suggestions:
            if isinstance(item, dict):
                _add_bullet(doc, item.get("suggestion") or item.get("action") or "")
    else:
        _add_body(doc, "暂无明显问题。")

    # Footer
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("由 ResumeMatch-AI 生成")
    set_run_font(rr, size=8)
    rr.font.color.rgb = RGBColor(150, 150, 150)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _add_score_table(doc, total_score: float, detail: ScoreDetail) -> None:
    from docx.shared import RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rows = [
        ("总分", f"{float(total_score):.1f}"),
        ("技能匹配", f"{detail.skill_score:.1f}"),
        ("项目相关性", f"{detail.project_score:.1f}"),
        ("语义相关度", f"{detail.semantic_score:.1f}"),
        ("硬性要求", f"{detail.requirement_score:.1f}"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = 1
        r = p.add_run(label + "\n")
        set_run_font(r, size=8, bold=True)
        r.font.color.rgb = RGBColor(90, 90, 90)
        r2 = p.add_run(value + " / 100")
        set_run_font(r2, size=12, bold=True)
        r2.font.color.rgb = RGBColor(31, 78, 121)
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F6FA")
        tc_pr.append(shd)


def _add_heading(doc, text: str) -> None:
    from docx.shared import RGBColor, Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(safe_export_text(text))
    set_run_font(r, size=12, bold=True)
    r.font.color.rgb = RGBColor(31, 78, 121)


def _add_subheading(doc, text: str) -> None:
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(safe_export_text(text))
    set_run_font(r, size=10.5, bold=True)


def _add_body(doc, text: str) -> None:
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(safe_export_text(text))
    set_run_font(r, size=10)


def _add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(safe_export_text(text))
    set_run_font(r, size=10)


def export_report_pdf(task_id: int, db: Session) -> bytes:
    """Generate a PDF report using ReportLab and a real CJK font.

    This avoids xhtml2pdf's common Chinese-font issue where unsupported glyphs
    become question marks.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except Exception as exc:
        raise AppException("缺少 reportlab，无法导出 PDF。请安装 requirements.txt。", code="PDF_EXPORT_FAILED") from exc

    task, report, resume, jd = _load_data(task_id, db)
    score_detail, ai = _load_score_detail(report)
    problems = _load_list(report.problems_json)
    matched = _load_list(report.matched_skills_json)
    missing = _load_list(report.missing_skills_json)
    suggestions = _load_list(report.suggestions_json)

    font = register_pdf_cjk_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=13 * mm, bottomMargin=13 * mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CNTitle", fontName=font, fontSize=18, leading=22, alignment=1, textColor=colors.HexColor("#1f4e79"), spaceAfter=6))
    styles.add(ParagraphStyle(name="CNMeta", fontName=font, fontSize=8.5, leading=11, alignment=1, textColor=colors.HexColor("#666666"), spaceAfter=8))
    styles.add(ParagraphStyle(name="CNHeading", fontName=font, fontSize=12, leading=15, textColor=colors.HexColor("#1f4e79"), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="CNBody", fontName=font, fontSize=9.3, leading=13.2, spaceAfter=3))
    styles.add(ParagraphStyle(name="CNBullet", fontName=font, fontSize=9.1, leading=12.5, leftIndent=10, firstLineIndent=-8, spaceAfter=2))

    story = []
    P = lambda text, style="CNBody": Paragraph(_html(safe_export_text(text)), styles[style])
    story.append(P("AI 简历 JD 匹配分析报告", "CNTitle"))
    story.append(P(f"简历：{resume.filename or ''}    岗位：{jd.title or ''}    公司：{jd.company or ''}", "CNMeta"))

    score_data = [[P("总分", "CNBody"), P("技能匹配", "CNBody"), P("项目相关", "CNBody"), P("语义相关", "CNBody"), P("硬性要求", "CNBody")],
                  [P(f"{float(report.total_score or 0):.1f}", "CNBody"), P(f"{score_detail.skill_score:.1f}", "CNBody"), P(f"{score_detail.project_score:.1f}", "CNBody"), P(f"{score_detail.semantic_score:.1f}", "CNBody"), P(f"{score_detail.requirement_score:.1f}", "CNBody")]]
    table = Table(score_data, colWidths=[34 * mm] * 5)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F6FA")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9E2EF")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([table, Spacer(1, 6)])

    story.append(P("综合结论", "CNHeading"))
    story.append(P(ai.get("overall_conclusion") or _build_basic_conclusion(report.total_score or 0, matched, missing)))

    story.append(P("评分说明", "CNHeading"))
    for label, score, reason in [
        ("技能匹配", score_detail.skill_score, score_detail.skill_reason),
        ("项目相关性", score_detail.project_score, score_detail.project_reason),
        ("语义相关度", score_detail.semantic_score, score_detail.semantic_reason),
        ("硬性要求", score_detail.requirement_score, score_detail.requirement_reason),
    ]:
        story.append(P(f"- {label}：{score:.1f} / 100。{reason}", "CNBullet"))

    story.append(P("技能匹配", "CNHeading"))
    story.append(P("命中技能：" + ("、".join(map(str, matched)) if matched else "暂无")))
    story.append(P("缺失/风险技能：" + ("、".join(map(str, missing)) if missing else "暂无")))

    story.append(P("简历问题与优化建议", "CNHeading"))
    if problems:
        for item in problems[:12]:
            if isinstance(item, dict):
                title = item.get("title") or "简历问题"
                story.append(P(f"- {title}：{item.get('suggestion') or ''}", "CNBullet"))
                evidence = item.get("evidence") or ""
                if evidence:
                    story.append(P(f"  证据：{evidence}", "CNBullet"))
    elif suggestions:
        for item in suggestions[:10]:
            if isinstance(item, dict):
                story.append(P("- " + str(item.get("suggestion") or item.get("action") or ""), "CNBullet"))
    else:
        story.append(P("暂无明显问题。"))

    if ai:
        if ai.get("resume_rewrite_strategy"):
            story.append(P("简历改写重点", "CNHeading"))
            for x in ai.get("resume_rewrite_strategy")[:8]:
                story.append(P("- " + str(x), "CNBullet"))
        if ai.get("interview_risks"):
            story.append(P("面试风险", "CNHeading"))
            for x in ai.get("interview_risks")[:8]:
                story.append(P("- " + str(x), "CNBullet"))

    doc.build(story)
    return buf.getvalue()


def export_report_html(task_id: int, db: Session) -> str:
    """Generate a simple HTML report for browser preview/debug only."""
    task, report, resume, jd = _load_data(task_id, db)
    score_detail, ai = _load_score_detail(report)
    matched = _load_list(report.matched_skills_json)
    missing = _load_list(report.missing_skills_json)
    conclusion = ai.get("overall_conclusion") or _build_basic_conclusion(report.total_score or 0, matched, missing)
    return f"""<!doctype html><html><head><meta charset='utf-8'><title>分析报告</title>
<style>body{{font-family:'Microsoft YaHei','SimSun',Arial,sans-serif;max-width:860px;margin:36px auto;line-height:1.65;color:#222}}h1,h2{{color:#1f4e79}}.score{{font-size:36px;color:#1f4e79;font-weight:700}}.card{{border:1px solid #e5e7eb;border-radius:10px;padding:16px;margin:12px 0;background:#fff}}</style></head><body>
<h1>AI 简历 JD 匹配分析报告</h1><p>简历：{html.escape(resume.filename or '')} | 岗位：{html.escape(jd.title or '')}</p>
<div class='card'><div class='score'>{float(report.total_score or 0):.1f} / 100</div><p>{html.escape(safe_export_text(conclusion))}</p></div>
<h2>评分详情</h2><ul><li>技能匹配：{score_detail.skill_score:.1f} - {html.escape(score_detail.skill_reason)}</li><li>项目相关：{score_detail.project_score:.1f} - {html.escape(score_detail.project_reason)}</li><li>语义相关：{score_detail.semantic_score:.1f} - {html.escape(score_detail.semantic_reason)}</li><li>硬性要求：{score_detail.requirement_score:.1f} - {html.escape(score_detail.requirement_reason)}</li></ul>
<h2>技能</h2><p>命中：{html.escape('、'.join(map(str, matched)) or '暂无')}</p><p>缺失：{html.escape('、'.join(map(str, missing)) or '暂无')}</p>
</body></html>"""


def _html(text: str) -> str:
    return html.escape(safe_export_text(text)).replace("\n", "<br/>")


def _db_get(db: Session, model: Any, pk: int):
    try:
        return db.get(model, pk)
    except AttributeError:
        return db.query(model).filter(model.id == pk).first()


def _load_data(task_id: int, db: Session) -> Tuple[AnalysisTask, MatchReport, Resume, JobDescription]:
    task = _db_get(db, AnalysisTask, task_id)
    if task is None:
        raise AppException("分析任务不存在", code="NOT_FOUND", status_code=404)
    report = db.query(MatchReport).filter(MatchReport.task_id == task_id).first()
    if report is None:
        raise AppException("分析报告不存在", code="NOT_FOUND", status_code=404)
    resume = _db_get(db, Resume, task.resume_id)
    jd = _db_get(db, JobDescription, task.jd_id)
    if resume is None or jd is None:
        raise AppException("简历或岗位不存在", code="NOT_FOUND", status_code=404)
    return task, report, resume, jd


def _load_list(raw: Any) -> list:
    if not raw:
        return []
    if isinstance(raw, list):
        return raw
    try:
        value = json.loads(raw)
        return value if isinstance(value, list) else []
    except Exception:
        return []


def _load_score_detail(report: MatchReport) -> tuple[ScoreDetail, dict[str, Any]]:
    raw: dict[str, Any] = {}
    if report.report_json:
        try:
            parsed = json.loads(report.report_json) if isinstance(report.report_json, str) else report.report_json
            raw = parsed if isinstance(parsed, dict) else {}
        except Exception:
            raw = {}
    ai = raw.get("ai_analysis") if isinstance(raw.get("ai_analysis"), dict) else {}
    score_payload = raw.get("score_detail") if isinstance(raw.get("score_detail"), dict) else raw
    try:
        detail = ScoreDetail(**(score_payload or {}))
    except Exception:
        detail = ScoreDetail()
    return detail, ai or {}


def _build_basic_conclusion(score: float, matched: list, missing: list) -> str:
    level = "强匹配" if score >= 85 else "推荐投递" if score >= 75 else "可投递，需优化" if score >= 60 else "低优先级"
    return f"综合匹配度 {float(score):.1f} 分，属于{level}。命中 {len(matched)} 项技能，存在 {len(missing)} 项待补充关键词。"
